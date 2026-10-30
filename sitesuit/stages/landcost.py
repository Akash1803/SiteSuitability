"""Stage 7 - tentative land cost in the team template (RFCTLARR style).
Condition 1 = MV 110 % of GLV + 12 % for 1 year; Condition 2 = 150 % + 12 % x 2 years. As per Act: MV x1.25 + 100 % solatium + 12 % x years;
private negotiation (GO 411): 275 % of MV + 25 % incentive; govt land = 50 % of the weighted patta compensation per acre. Sq-ft rates x 43,560."""
import re

import numpy as np
import openpyxl
import pandas as pd

from .. import db
from ..config import SETTINGS
from ..util import GOVT, HA_TO_ACRE, dc_keys, glv_final

NAME, TITLE = "landcost", "Tentative land cost"


def _svs(s):
    return ", ".join(sorted(set(s), key=lambda x: float(re.sub(r"\D", "", x) or 0)))


def _rates(sm):
    """per-record rate per acre for the template (sq-ft x 43,560); patta records without a usable rate take the village dry rate"""
    sm = sm.copy()
    sm["own"] = np.where(sm.government_priority == GOVT, "Govt", "Patta")
    sm["rate"] = pd.to_numeric(sm.glv_rate, errors="coerce") * np.where(sm.glv_unit.astype(str) == "Square Feet", 43560.0, 1.0)
    sm["basis"] = sm.glv_match.astype(str).str.split(" (", regex=False).str[0].str.split(" - ").str[0]
    pat = sm[(sm.own == "Patta") & (sm.rate > 0) & (sm.glv_unit.astype(str) == "Acre")]
    vill_dry = pat.groupby("village").rate.agg(lambda s: s.value_counts().index[0]).to_dict()
    any_rate = float(pat.rate.mode().iloc[0]) if len(pat) else 0.0
    assumed = []
    for i in sm.index[(sm.own == "Patta") & ((sm.rate.isna()) | (sm.rate == 0))]:
        r = vill_dry.get(sm.at[i, "village"], any_rate)
        sm.at[i, "rate"] = r; sm.at[i, "basis"] = "Assumed village dry rate (register 0 / no rate)"
        assumed.append((sm.at[i, "unique_id"], sm.at[i, "site_acre"], r))
    return sm, assumed


def _cost(rate, pct, yrs):
    """compensation per acre: (as per Act, by private negotiation)"""
    G = rate * pct / 100
    return G * 1.25 * 2 + G * 12 * yrs / 100, G * 2.75 * 1.25


def dc_split(sm, dc):
    """the land cost split into Non-DC / DC land (dc = boolean Series on sm's index). Same formulas as compute(); govt land in both
    groups at 50% of the whole site's weighted patta compensation per acre, so the two groups add up to the site totals."""
    sm, _ = _rates(sm)
    sm["dc"] = dc.reindex(sm.index).fillna(False).astype(bool)
    pat, gov = sm[sm.own == "Patta"], sm[sm.own == "Govt"]
    PA = float(pat.site_acre.sum())
    out = {}
    for cond, pct, yrs in (("C1", 110, 1), ("C2", 150, 2)):
        ka, kn = _cost(pat.rate, pct, yrs)
        act, neg = ka * pat.site_acre, kn * pat.site_acre
        wa_a, wa_n = (act.sum() / PA, neg.sum() / PA) if PA else (0, 0)
        res = {}
        for grp, m in (("non_dc", ~sm.dc), ("dc", sm.dc), ("total", pd.Series(True, index=sm.index))):
            mp, mg = m.reindex(pat.index), m.reindex(gov.index)
            pa, ga = float(pat[mp].site_acre.sum()), float(gov[mg].site_acre.sum())
            r = {"gov": {"acre": ga, "act": 0.5 * wa_a * ga, "neg": 0.5 * wa_n * ga}, "patta": {"acre": pa, "act": float(act[mp].sum()), "neg": float(neg[mp].sum())}}
            t = {k: r["gov"][k] + r["patta"][k] for k in ("acre", "act", "neg")}
            t["act_per_acre"], t["neg_per_acre"] = (t["act"] / t["acre"], t["neg"] / t["acre"]) if t["acre"] else (0, 0)
            r["total"] = t
            res[grp] = r
        out[cond] = res
    return out


def compute(sm):
    sm, assumed = _rates(sm)
    pat = sm[sm.own == "Patta"]; gov = sm[sm.own == "Govt"]
    groups = pat.groupby(["village", "rate", "glv_unit", "basis"], dropna=False).agg(acre=("site_acre", "sum"), n=("unique_id", "size"), surveys=("survey_no", _svs)).reset_index()
    groups = groups.sort_values(["village", "rate"]).reset_index(drop=True)
    groups["unit_label"] = np.where(groups.glv_unit.astype(str) == "Square Feet", (groups.rate / 43560).round(0).astype(int).astype(str) + "/ Square Feet", groups.rate.astype(int).astype(str) + "/ Acre")
    gov_rows = gov.groupby(["village", "land_type_english"]).agg(acre=("site_acre", "sum"), surveys=("survey_no", _svs)).reset_index()
    res = {}
    for cond, pct, yrs in (("C1", 110, 1), ("C2", 150, 2)):
        G = groups.rate.values * pct / 100; H = G * 1.25; K = H + H + G * 12 * yrs / 100; M = G * 275 / 100; O = M + M * 0.25
        pa, ga = groups.acre.sum(), gov.site_acre.sum()
        act_p, neg_p = (K * groups.acre.values).sum(), (O * groups.acre.values).sum()
        wa_a, wa_n = (act_p / pa, neg_p / pa) if pa else (0, 0)
        act_g, neg_g = 0.5 * wa_a * ga, 0.5 * wa_n * ga
        t = pa + ga
        res[cond] = dict(pct=pct, years=yrs, pat_acre=pa, gov_acre=ga, total_acre=t, act_pat=act_p, neg_pat=neg_p, act_gov=act_g, neg_gov=neg_g,
                         act_total=act_p + act_g, neg_total=neg_p + neg_g, act_per_acre=(act_p + act_g) / t if t else 0, neg_per_acre=(neg_p + neg_g) / t if t else 0)
    return groups, gov_rows, res, assumed


def write_workbook(site, template, groups, gov_rows, res, assumed):
    wb = openpyxl.load_workbook(template)
    sp, sg, sf = wb["Patta"], wb["Poromboke"], wb["Final_Calculation"]
    title_site = f"{site.name} SIPCOT site"

    def fill(wp, wg, wf, pct, yrs, title):
        for r in range(5, 1843):
            for c in range(1, 17):
                wp.cell(r, c).value = None
        for k, g in groups.iterrows():
            r = 5 + k
            wp.cell(r, 1).value = k + 1; wp.cell(r, 2).value = g.village; wp.cell(r, 3).value = f"Survey {g.surveys} ({int(g.n)} sub-div; {g.basis}; GLV {g.unit_label})"
            wp.cell(r, 4).value = round(g.acre / HA_TO_ACRE, 4); wp.cell(r, 5).value = round(g.acre, 4); wp.cell(r, 6).value = float(g.rate)
            wp.cell(r, 7).value = f"=(F{r}*{pct}/100)"; wp.cell(r, 8).value = f"=G{r}*1.25"; wp.cell(r, 9).value = f"=H{r}"; wp.cell(r, 10).value = f"=(G{r}*12*{yrs})/100"
            wp.cell(r, 11).value = f"=H{r}+I{r}+J{r}"; wp.cell(r, 12).value = f"=K{r}*E{r}"; wp.cell(r, 13).value = f"=G{r}*275/100"; wp.cell(r, 14).value = f"=M{r}*25/100"
            wp.cell(r, 15).value = f"=M{r}+N{r}"; wp.cell(r, 16).value = f"=O{r}*E{r}"
        last = 5 + max(len(groups), 1) - 1
        wp.cell(1843, 5).value = f"=SUM(E5:E{last})"; wp.cell(1843, 12).value = f"=SUM(L5:L{last})"; wp.cell(1843, 16).value = f"=SUM(P5:P{last})"
        wp.cell(1, 1).value = f"TENTATIVE LAND COST (Patta Land) - {title_site}, {title}"
        wp.cell(3, 7).value = f"Market Value ({pct}% of GLV) (in Rs.)"; wp.cell(3, 10).value = f"12% additional MV p.a for {yrs} Year{'s' if yrs > 1 else ''} (on Col.7) (in Rs.)"
        for r in range(3, 76):
            for c in range(1, 12):
                wg.cell(r, c).value = None
        for k, g in gov_rows.iterrows():
            r = 3 + k
            wg.cell(r, 1).value = k + 1; wg.cell(r, 2).value = g.village; wg.cell(r, 3).value = f"Govt {g.land_type_english}: Survey {g.surveys}"
            wg.cell(r, 4).value = round(g.acre / HA_TO_ACRE, 4); wg.cell(r, 5).value = round(g.acre, 4)
            wg.cell(r, 6).value = f"={wp.title}!L1843/{wp.title}!E1843"; wg.cell(r, 7).value = f"=F{r}*0.5"; wg.cell(r, 8).value = f"=G{r}*E{r}"
            wg.cell(r, 9).value = f"={wp.title}!P1843/{wp.title}!E1843"; wg.cell(r, 10).value = f"=I{r}*0.5"; wg.cell(r, 11).value = f"=J{r}*E{r}"
        gl = 3 + max(len(gov_rows), 1) - 1
        wg.cell(77, 5).value = f"=SUM(E3:E{gl})"; wg.cell(77, 8).value = f"=SUM(H3:H{gl})"; wg.cell(77, 11).value = f"=SUM(K3:K{gl})"
        tot = round(groups.acre.sum() + gov_rows.acre.sum(), 2)
        wf["E3"].value = f"{title_site} - Land allotment at {pct}% Market Value with 12% additional MV p.a. for {yrs} year{'s' if yrs > 1 else ''} ({title})"
        wf["G5"].value = f"={wg.title}!E77"; wf["H5"].value = f"={wg.title}!H77"; wf["I5"].value = f"={wg.title}!K77"
        wf["G6"].value = f"={wp.title}!E1843"; wf["H6"].value = f"={wp.title}!L1843"; wf["I6"].value = f"={wp.title}!P1843"
        wf["E7"].value = f"Total Proposed Land: {tot} Acres (A-register); Compensation Grand Total (Rupees)"
        wf["H7"].value = "=SUM(H5:H6)"; wf["I7"].value = "=SUM(I5:I6)"; wf["H8"].value = None; wf["I8"].value = None
        wf["H9"].value = f"=H7/{tot}"; wf["I9"].value = f"=I7/{tot}"; wf["E9"].value = "Estimated cost per acre (Rs.)"; wf["I13"].value = None

    p2 = wb.copy_worksheet(sp); p2.title = "Patta_150pct_2yr"
    g2 = wb.copy_worksheet(sg); g2.title = "Poromboke_150pct_2yr"
    f2 = wb.copy_worksheet(sf); f2.title = "Final_150pct_2yr"
    sp.title = "Patta_110pct_1yr"; sg.title = "Poromboke_110pct_1yr"; sf.title = "Final_110pct_1yr"
    fill(sp, sg, sf, 110, 1, "Condition 1"); fill(p2, g2, f2, 150, 2, "Condition 2")
    ws = wb.create_sheet("GLV_Summary", 0)
    c1, c2 = res["C1"], res["C2"]
    for row in [[f"{title_site} - Guideline Land Value (GLV) based tentative land cost"],
                [f"Site area (A-register; split edge surveys at the area inside): patta {c1['pat_acre']:.2f} ac + govt {c1['gov_acre']:.2f} ac = {c1['total_acre']:.2f} ac"], [],
                ["S.No", "Land Type", "Acres", "Condition 1 - Total compensation as per Act", "Condition 1 - as per private negotiation", "Condition 2 - Total compensation as per Act", "Condition 2 - as per private negotiation"],
                [1, "GOV LAND", round(c1["gov_acre"], 3), round(c1["act_gov"]), round(c1["neg_gov"]), round(c2["act_gov"]), round(c2["neg_gov"])],
                [2, "PRIVATE LAND", round(c1["pat_acre"], 3), round(c1["act_pat"]), round(c1["neg_pat"]), round(c2["act_pat"]), round(c2["neg_pat"])],
                ["", f"Total {c1['total_acre']:.2f} acres - Compensation grand total (Rs)", round(c1["total_acre"], 3), round(c1["act_total"]), round(c1["neg_total"]), round(c2["act_total"]), round(c2["neg_total"])],
                ["", "Estimated cost per acre (Rs)", "", round(c1["act_per_acre"]), round(c1["neg_per_acre"]), round(c2["act_per_acre"]), round(c2["neg_per_acre"])],
                ["", "Estimated cost per acre (Rs crore)", "", round(c1["act_per_acre"] / 1e7, 2), round(c1["neg_per_acre"] / 1e7, 2), round(c2["act_per_acre"] / 1e7, 2), round(c2["neg_per_acre"] / 1e7, 2)], [],
                ["Condition 1: Market value = 110% of GLV, 12% additional MV p.a. for 1 year.  Condition 2: Market value = 150% of GLV, 12% additional MV p.a. for 2 years."],
                ["As per Act: MV x 1.25 + 100% solatium + 12% x years on MV; Private negotiation (GO 411): 275% of MV + 25% incentive. Govt land: 50% of the weighted-average patta compensation per acre."],
                ["GLV source: TNReginet survey-wise entries as quoted (sq-ft for layouts, acre for land); unlisted subdivisions take the lowest rate of their survey. Sq-ft rates x 43,560 per acre as in the template."],
                ["Numbers here are static; the Patta_/Poromboke_/Final_ sheets hold the live template formulas."]]:
        ws.append(row)
    rb = wb.create_sheet("Rate_Basis", 1)
    rb.append(["Village", "Rate entered in template (Rs/acre; sq-ft rates x 43,560)", "GLV as quoted", "Basis", "Sub-divisions", "Acres (site)", "Surveys"])
    for _, g in groups.iterrows():
        rb.append([g.village, float(g.rate), g.unit_label, g.basis, int(g.n), round(g.acre, 4), g.surveys])
    rb.append([]); rb.append(["Govt land rows:"])
    for _, g in gov_rows.iterrows():
        rb.append([g.village, "", "", f"Govt {g.land_type_english} (50% of weighted patta compensation)", "", round(g.acre, 4), g.surveys])
    if assumed:
        rb.append([]); rb.append(["Patta records with no usable register rate - village dry rate assumed:"])
        for u, a, r in assumed:
            rb.append([u, float(r), "", "assumed", "", round(float(a), 4), ""])
    wb.save(site.landcost)


def stale(site, ctx):
    """the DC-land list (01_Source/DC_Land_Surveys.xlsx) was added or changed after the cost was worked out"""
    summ = db.get_summary(site.site_id)
    fin = glv_final(site)
    used = (summ.get("landcost") or {}).get("final_source")
    return (sorted(dc_keys(site)) != (summ.get("dc_list", summ.get("dc_surveys")) or [])
            or bool(fin) != bool(used) or (fin and fin.get("source") != used))


def run(site, ctx):
    sm = pd.read_excel(site.master, sheet_name="Subdivision_Master", dtype={"survey_no": str, "subdiv_no": str})
    sm = sm[sm.in_site.astype(str).str.lower() == "true"].copy()
    groups, gov_rows, res, assumed = compute(sm)
    write_workbook(site, SETTINGS.landcost_template, groups, gov_rows, res, assumed)
    db.register_document(site.site_id, "landcost", site.landcost)
    keep = {k: {kk: float(vv) for kk, vv in v.items()} for k, v in res.items()}
    fin = glv_final(site)
    if fin:   # the final GLV calculation was supplied (feasibility platform): the slides use it; the workbook keeps this app's working
        for k in ("C1", "C2"):
            c = fin[k]
            keep[k].update(pat_acre=c["patta"]["acre"], gov_acre=c["gov"]["acre"], total_acre=c["total"]["acre"],
                           act_pat=c["patta"]["act"] * 1e7, neg_pat=c["patta"]["neg"] * 1e7, act_gov=c["gov"]["act"] * 1e7, neg_gov=c["gov"]["neg"] * 1e7,
                           act_total=c["total"]["act"] * 1e7, neg_total=c["total"]["neg"] * 1e7,
                           act_per_acre=c["total"]["act"] * 1e7 / c["total"]["acre"], neg_per_acre=c["total"]["neg"] * 1e7 / c["total"]["acre"])
        keep["final_source"] = fin.get("source", "supplied final GLV figures")
        ctx.log(f"GLV final figures taken from {fin.get('source', 'the supplied file')}: C1 {fin['C1']['total']['act']:.2f} / {fin['C1']['total']['neg']:.2f} cr, "
                f"C2 {fin['C2']['total']['act']:.2f} / {fin['C2']['total']['neg']:.2f} cr on {fin['C1']['total']['acre']:.2f} ac")
        res = {k: keep[k] for k in ("C1", "C2")}
    dcs = dc_keys(site)
    split = dc_split(sm, sm.survey_uid.astype(str).isin(dcs)) if dcs else None
    if split and split["C1"]["dc"]["total"]["acre"] <= 0:
        ctx.log(f"DC land list has {len(dcs)} surveys but none of them is in the site - the GLV summary stays without the DC split"); split = None
    db.save_summary(site.site_id, {"landcost": keep, "landcost_assumed": len(assumed), "landcost_dc": split,
                                   "dc_surveys": sorted(dcs) if split else [], "dc_list": sorted(dcs)})
    c1, c2 = res["C1"], res["C2"]
    if split:
        d1 = split["C1"]["dc"]["total"]
        ctx.log(f"DC land: {d1['acre']:.2f} ac, C1 {d1['act'] / 1e7:.2f} cr as per Act ({d1['act_per_acre'] / 1e5:.2f} lakh/ac) - the GLV summary slide is split into DC / Non-DC land")
    return (f"C1: {c1['act_per_acre'] / 1e5:.2f} lakh/ac as per Act, {c1['neg_per_acre'] / 1e5:.2f} by negotiation; "
            f"C2: {c2['act_per_acre'] / 1e5:.2f} / {c2['neg_per_acre'] / 1e5:.2f} lakh/ac" + (f"; {len(assumed)} patta records on an assumed rate" if assumed else ""))
