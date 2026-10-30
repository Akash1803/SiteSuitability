"""Stage 10 - slide images (map slides composed with PIL, summary slides from HTML) and the PPT built from the deck template.
No PDF here: the PDF is made only after the PPT is confirmed (button in the app)."""
import json
import math
import shutil
import subprocess
import sys
from pathlib import Path

import pandas as pd
from PIL import Image, ImageDraw, ImageFont

from .. import db
from ..config import SETTINGS
from ..pin import teardrop
from ..util import glv_final

NAME, TITLE = "ppt", "Slides + PPT"
W, H, B = 2160, 1374, 126
F = lambda s: ImageFont.truetype("C:/Windows/Fonts/arialbd.ttf", s)


def _and(xs):
    xs = list(dict.fromkeys(xs))
    return xs[0] if len(xs) == 1 else ", ".join(xs[:-1]) + " and " + xs[-1]


def to3857(lon, lat):
    return lon * 20037508.34 / 180, math.log(math.tan((90 + lat) * math.pi / 360)) * 6378137


LINK_BOX = (1270, 1300, 2160, 1500)   # template's "view on Google Maps" text box (boundary slide), in slide pixels
TOOLBAR = (0, 1430, 760, 1500)        # PowerPoint slide-show buttons (bottom-left)


def _overlap(a, b):
    return max(0, min(a[2], b[2]) - max(a[0], b[0])) * max(0, min(a[3], b[3]) - max(a[1], b[1]))


def best_corner(pw, ph, avoid, margin=24, bottom_gap=24):
    """the corner box (x0, y0) that covers the least of the boxes in avoid"""
    cands = [(W - pw - margin, B + margin), (margin, B + margin), (margin, B + H - ph - bottom_gap), (W - pw - margin, B + H - ph - bottom_gap)]
    return min(cands, key=lambda c: sum(_overlap((c[0], c[1], c[0] + pw, c[1] + ph), a) * wgt for a, wgt in avoid))


def site_box(site, e):
    import geopandas as gpd
    g = gpd.read_file(site.boundary).to_crs(3857).geometry.iloc[0]
    x0, y0, x1, y1 = g.bounds
    px = lambda x, y: ((x - e[0]) / (e[2] - e[0]) * W, B + (e[3] - y) / (e[3] - e[1]) * H)
    a, b_ = px(x0, y1), px(x1, y0)
    return (a[0], a[1], b_[0], b_[1])


def compose(site, raw, out, AC, summ, cn):
    I = json.loads((raw / "_render_info.json").read_text())

    def base(name, title):
        im = Image.open(raw / name).convert("RGB"); o = Image.new("RGB", (W, H + B), (16, 58, 78)); o.paste(im, (0, B)); d = ImageDraw.Draw(o)
        f = F(66); d.text(((W - d.textlength(title, font=f)) / 2, (B - 76) / 2), title, fill=(255, 235, 0), font=f); return o, d

    def pxe(e, x, y): return (x - e[0]) / (e[2] - e[0]) * W, B + (e[3] - y) / (e[3] - e[1]) * H

    def label(d, xy, tx, size, col=(255, 235, 0)):
        f = F(size); d.text((xy[0] - d.textlength(tx, font=f) / 2, xy[1] - size * 0.6), tx, fill=col, font=f, stroke_width=4, stroke_fill=(0, 0, 0))

    def panel(d, x0, y0, title, lines, fs, right=False, bottom=False, avoid=None):
        fp = F(fs); pw = max([d.textlength(s, font=fp) for s in lines] + [d.textlength(title, font=F(36))]) + 48; ph = len(lines) * (fs + 16) + 76
        if right: x0 = W - pw - 24
        if bottom: y0 = B + H - ph - 20
        if avoid is not None: x0, y0 = best_corner(pw, ph, avoid, bottom_gap=80)
        d.rectangle([x0, y0, x0 + pw, y0 + ph], fill=(16, 58, 78), outline=(255, 255, 255), width=3); d.text((x0 + 24, y0 + 14), title, fill=(255, 235, 0), font=F(36))
        for i, s in enumerate(lines): d.text((x0 + 24, y0 + 68 + i * (fs + 16)), s, fill=(255, 255, 255), font=fp)
        return x0, y0, pw, ph

    # boundary
    o, d = base("boundary_raw.png", "BOUNDARY"); e = I["boundary"]
    dm = summ.get("dimensions")
    if dm:
        FC = [(255, 61, 0), (0, 229, 255), (255, 214, 0), (213, 0, 249), (118, 255, 3)]
        for k, f in enumerate(dm["fronts"]):
            col = FC[k % len(FC)]
            lines = f["geom"]["coordinates"] if f["geom"]["type"] == "MultiLineString" else [f["geom"]["coordinates"]]
            for ln in lines:
                pts = [pxe(e, *to3857(x, y)) for x, y in ln]
                if len(pts) > 1:
                    d.line(pts, fill=col, width=12)
            mid = [pxe(e, *to3857(x, y)) for x, y in lines[0]][len(lines[0]) // 2]
            t = f"{f['label']} frontage: {f['m']:,} m"
            d.text((min(max(mid[0] + 16, 10), W - d.textlength(t, font=F(30)) - 10), min(max(mid[1] - 46, B + 10), B + H - 50)), t, fill=col, font=F(30), stroke_width=4, stroke_fill=(0, 0, 0))
        for key, txt in (("long_line", f"{dm.get('long_span_m', dm['length_m']) / 1000:.2f} km"), ("short_line", f"{dm.get('short_span_m', dm['width_m']) / 1000:.2f} km")):
            (a1, a2) = [pxe(e, *to3857(x, y)) for x, y in dm[key]]
            d.line([a1, a2], fill=(255, 255, 255), width=4)
            for pa, pb in ((a1, a2), (a2, a1)):
                ang = math.atan2(pb[1] - pa[1], pb[0] - pa[0])
                for sgn in (1, -1):
                    d.line([pa, (pa[0] + 30 * math.cos(ang + sgn * 0.45), pa[1] + 30 * math.sin(ang + sgn * 0.45))], fill=(255, 255, 255), width=4)
            mx, my = a1[0] + (a2[0] - a1[0]) * 0.72, a1[1] + (a2[1] - a1[1]) * 0.72   # label near one end, clear of the acreage label
            d.text((mx + 14, my - 40), txt, fill=(255, 255, 255), font=F(32), stroke_width=4, stroke_fill=(0, 0, 0))
        info = [f"Area: {AC} acres", f"Overall extent: {dm['length_m'] / 1000:.2f} km x {dm['width_m'] / 1000:.2f} km", f"Perimeter: {dm['perimeter_m'] / 1000:.2f} km"]
        info += [f"{f['label']} frontage: {f['m']:,} m" for f in dm["fronts"][:4]] or ["No frontage on a named road"]
        panel(d, 0, 0, "Site Dimensions", info, 27, avoid=[(site_box(site, e), 1.0), (LINK_BOX, 3.0), (TOOLBAR, 2.0)])
    label(d, pxe(e, *I["site_pole"]), AC + " Acres", 54); o.save(out / "Boundary.png")
    # road connectivity
    o, d = base("road_connectivity_raw.png", "ROAD CONNECTIVITY"); e = I["road"]
    qx, qy = pxe(e, *I["site_pole"])
    for i, s in enumerate(("Site Boundary", f"({AC} Acres)")):
        d.text((qx - d.textlength(s, font=F(34)) / 2, qy - 44 + i * 44), s, fill=(255, 255, 255), font=F(34), stroke_width=4, stroke_fill=(0, 0, 0))
    for nm, lon, lat, _ in cn.get("stations", []):
        x, y = pxe(e, *to3857(lon, lat))
        if 0 < x < W and B < y < B + H:
            d.rectangle([x - 13, y - 13, x + 13, y + 13], fill=(0, 229, 255), outline=(0, 0, 0), width=3)
            t = f"{nm} Rly Stn"; tw = d.textlength(t, font=F(30))
            d.text((x + 20 if x + 20 + tw < W - 10 else x - 20 - tw, y - 18), t, fill=(0, 229, 255), font=F(30), stroke_width=4, stroke_fill=(0, 0, 0))
    for bear, txt in cn.get("arrows", []):
        cx, cy = W / 2, B + H / 2; rr = min(W, H) / 2 - 70; aa = math.radians(bear)
        ex = min(max(cx + rr * 1.15 * math.sin(aa), 70), W - 70); ey = min(max(cy - rr * math.cos(aa), B + 70), B + H - 70)
        x0, y0 = ex - 200 * math.sin(aa), ey + 200 * math.cos(aa); d.line([(x0, y0), (ex, ey)], fill=(255, 255, 255), width=8)
        for s in (150, -150):
            d.line([(ex, ey), (ex + 34 * math.sin(aa + math.radians(s)), ey - 34 * math.cos(aa + math.radians(s)))], fill=(255, 255, 255), width=8)
        tx = min(max(x0 - 140, 12), W - 430); ty = min(max(y0 + (24 if bear < 90 or bear > 270 else -110), B + 12), B + H - 110)
        d.multiline_text((tx, ty), txt, fill=(255, 255, 0), font=F(34), stroke_width=4, stroke_fill=(0, 0, 0))
    px0, py0, ppw, pph = panel(d, 20, 0, f"Site Extent: {AC} Acres", ["• " + s for s in cn["lines"][:9]], 25, avoid=[(site_box(site, e), 1.0), (TOOLBAR, 2.0)])
    corners_l = [(W - 498, B + H - 196 - 60), (24, B + H - 196 - 60), (W - 498, B + 24), (24, B + 24)]
    lx, ly = min(corners_l, key=lambda c: _overlap((c[0], c[1], c[0] + 470, c[1] + 168), (px0, py0, px0 + ppw, py0 + pph)) * 5
                 + _overlap((c[0], c[1], c[0] + 470, c[1] + 168), site_box(site, e)))
    d.rectangle([lx, ly, lx + 470, ly + 168], fill=(16, 58, 78), outline=(255, 255, 255), width=2)
    for i, (c, t) in enumerate([((212, 255, 0), "Major district road (MDR)"), ((255, 26, 26), "Village road"), ((255, 255, 255), "Railway line / site boundary")]):
        d.line([(lx + 18, ly + 32 + i * 46), (lx + 74, ly + 32 + i * 46)], fill=c, width=7); d.text((lx + 90, ly + 16 + i * 46), t, fill=(255, 255, 255), font=F(27))
    o.save(out / "Road_Connectivity.png")
    # location
    R = I["region"]; e = R["ext"]
    o, d = base("region_raw.png", "LOCATION MAP")
    for n, (x, y) in R["dlab"].items():
        px, py = pxe(e, x, y); t = n.upper() + " DISTRICT"; hw = d.textlength(t, font=F(40)) / 2
        label(d, (min(max(px, hw + 20), W - hw - 20), py), t, 40, (120, 210, 255))
    x, y = pxe(e, *R["site_c"])
    pin, (tx, ty) = teardrop(r=36, h=124); o.paste(pin, (int(round(x - tx)), int(round(y - ty))), pin)
    d = ImageDraw.Draw(o)
    vs = [v["village"] for v in summ["villages"]]; tl = [v["taluk"] for v in summ["villages"]]; ds = [v["district"] for v in summ["villages"]]
    vl = ", ".join(vs); lines = []
    while vl:   # wrap the village list at ~44 characters
        cut = vl[:44].rfind(", ") if len(vl) > 44 else len(vl)
        cut = len(vl) if cut <= 0 else cut
        lines.append(("Villages: " if not lines else "   ") + vl[:cut]); vl = vl[cut + 2:] if cut < len(vl) else ""
    lines += [f"Taluk{'s' if len(set(tl)) > 1 else ''}: {_and(tl)}", f"District{'s' if len(set(ds)) > 1 else ''}: {_and(ds)}"] + cn["lines"][:5]
    if cn.get("nh") and cn["nh"] not in lines:
        lines.append(cn["nh"])
    if cn.get("sipcot_park"):
        lines.append(cn["sipcot_park"])
    px0, py0, pw, ph = panel(d, 0, B + 24, "Site Details", lines, 27, right=True)
    lt = f"Proposed Site ({AC} Acres)"; hw = d.textlength(lt, font=F(38)) / 2; ly = y - 158
    lx = x if not (x + hw > px0 - 10 and ly < py0 + ph) else px0 - 10 - hw
    label(d, (max(lx, hw + 10), ly), lt, 38)
    x0, y0 = 28, B + H - 150; d.rectangle([x0, y0, x0 + 380, y0 + 122], fill=(16, 58, 78), outline=(255, 255, 255), width=2)
    for i, (cc, t) in enumerate([((63, 182, 255), "District boundary"), ((255, 255, 255), "Taluk boundary")]):
        d.line([(x0 + 20, y0 + 38 + i * 46), (x0 + 80, y0 + 38 + i * 46)], fill=cc, width=6); d.text((x0 + 96, y0 + 20 + i * 46), t, fill=(255, 255, 255), font=F(28))
    o.save(out / "Location_Map.png")


def inr(v):
    """Indian digit grouping: 220000 -> '₹2,20,000'"""
    n = f"{int(round(float(v)))}"
    if len(n) <= 3:
        return "₹" + n
    head, groups = n[:-3], []
    while len(head) > 2:
        groups.insert(0, head[-2:]); head = head[:-2]
    return "₹" + ",".join(([head] if head else []) + groups + [n[-3:]])


def _ranges(nums):
    """[499, 501, 502, 503, 506] -> '499, 501–503, 506'"""
    out, run = [], []
    for n in sorted(nums):
        if run and n == run[-1] + 1:
            run.append(n); continue
        if run:
            out.append(f"{run[0]}–{run[-1]}" if len(run) > 2 else ", ".join(map(str, run)))
        run = [n]
    if run:
        out.append(f"{run[0]}–{run[-1]}" if len(run) > 2 else ", ".join(map(str, run)))
    return ", ".join(out)


def _dc_glv(s, pa, summ):
    """data for the GLV summary split into DC / Non-DC land (only when the site has a DC-land list, see util.dc_keys)"""
    split, dcs = summ.get("landcost_dc"), set(summ.get("dc_surveys") or [])
    if not split or not dcs:
        return None
    isdc = s.survey_uid.astype(str).isin(dcs)
    groups = (("non_dc", ~isdc), ("dc", isdc), ("total", isdc | ~isdc))
    rs = inr
    units = {}
    acre_m = (pa.glv_unit == "Acre") & (pa.glv_rate > 0); sq_m = pa.glv_unit == "Square Feet"
    for name, m in (("Acres", acre_m), ("Square feet", sq_m), ("No data", ~acre_m & ~sq_m)):
        x = pa[m]
        u = {g: float(x[mm.reindex(x.index)].site_acre.sum()) for g, mm in groups}
        if name == "No data" and u["total"] <= 0:
            continue
        u["range"] = {}
        for g, mm in groups[:2]:
            y = x[mm.reindex(x.index)].glv_rate
            u["range"][g] = "–" if name == "No data" or not len(y) else (rs(y.min()) if y.min() == y.max() else f"{rs(y.min())} – {rs(y.max())}")
        units[name] = u
    by_v = {}
    for k in sorted(dcs):
        hit = s[s.survey_uid.astype(str) == k]
        if len(hit):
            by_v.setdefault(hit.village.iloc[0], []).append(str(hit.survey_no.iloc[0]))
    txt = "; ".join(f"{v} survey nos. " + (_ranges([int(x) for x in sv]) if all(x.isdigit() for x in sv) else ", ".join(sv)) for v, sv in by_v.items())
    return {"split": split, "units": units, "patta_acre": float(pa.site_acre.sum()), "surveys_text": txt, "n_surveys": sum(len(v) for v in by_v.values())}


def slide_data(site, summ):
    s = pd.read_excel(site.master, sheet_name="Subdivision_Master"); s = s[s.in_site.astype(str).str.lower() == "true"]
    A = float(s.site_acre.sum())
    row = lambda v: {"total": round(float(v), 2), "pct": round(float(v) / A * 100, 2)}
    gov = s[s.ownership == "Govt"].site_acre.sum(); pat = s[s.ownership == "Patta"].site_acre.sum()
    unk = float(s[s.ownership.astype(str).str.startswith("Unknown")].site_acre.sum())   # the register has no record for these
    names = {"Dry": "Dry (Punjai)", "Wet": "Wet (Nanjai)", "Poramboke": "Poramboke (Government)", "Tharisu": "Tharisu (Waste land)", "Manavari": "Manavari (Rain-fed)"}
    g = s[s.ownership == "Govt"]
    rem = g.groupby(g.remarks.fillna("-").replace("-", "No remarks")).site_acre.sum().sort_values(ascending=False)
    pa = s[s.ownership == "Patta"]

    def rng(u):
        v = pa[(pa.glv_unit == u) & (pa.glv_rate > 0)].glv_rate
        return (f"₹{v.min():,.0f} – ₹{v.max():,.0f}" if len(v) else "–"), float(pa[(pa.glv_unit == u) & (pa.glv_rate > 0)].site_acre.sum())
    ac_r, ac_a = rng("Acre"); sq_r, sq_a = rng("Square Feet")
    nod = float(pa.site_acre.sum()) - ac_a - sq_a
    lc = summ["landcost"]

    def cond(k):
        c = lc[k]; t = c["total_acre"]
        return {"gov": {"acre": c["gov_acre"], "pct": c["gov_acre"] / t * 100, "act": c["act_gov"] / 1e7, "neg": c["neg_gov"] / 1e7},
                "patta": {"acre": c["pat_acre"], "pct": c["pat_acre"] / t * 100, "act": c["act_pat"] / 1e7, "neg": c["neg_pat"] / 1e7},
                "total": {"acre": c["total_acre"] if lc.get("final_source") else A, "act": c["act_total"] / 1e7, "neg": c["neg_total"] / 1e7},
                "per_acre_act_lakh": c["act_per_acre"] / 1e5, "per_acre_neg_lakh": c["neg_per_acre"] / 1e5}
    crop = summ.get("crop") or {"crop_acre": 0.0, "no_survey_acre": A, "noncrop_acre": 0.0}
    vil = s.groupby("village").site_acre.sum()
    # land abstract: village x (government / patta) x land type
    lt = s.land_type_english.fillna("Other")
    gtypes = [t for t in ("Dry", "Wet", "Poramboke", "Tharisu") if ((s.ownership == "Govt") & (lt == t)).any()] + sorted(set(lt[s.ownership == "Govt"]) - {"Dry", "Wet", "Poramboke", "Tharisu"})
    ptypes = [t for t in ("Dry", "Wet") if ((s.ownership == "Patta") & (lt == t)).any()] + sorted(set(lt[s.ownership == "Patta"]) - {"Dry", "Wet"})
    def _ab(x, name):
        return {"village": name, "govt": {t: float(x[(x.ownership == "Govt") & (x.land_type_english.fillna("Other") == t)].site_acre.sum()) for t in gtypes},
                "patta": {t: float(x[(x.ownership == "Patta") & (x.land_type_english.fillna("Other") == t)].site_acre.sum()) for t in ptypes},
                "unknown": float(x[x.ownership.astype(str).str.startswith("Unknown")].site_acre.sum())}
    abstract = {"govt_types": gtypes, "patta_types": ptypes, "rows": [_ab(x, v) for v, x in s.groupby("village")], "total": _ab(s, "Grand Total")}
    oc = []
    for t, c in (("Dry", "#f5c242"), ("Wet", "#4a8df0")):
        v = float(s[(s.ownership == "Patta") & (lt == t)].site_acre.sum())
        if v > 0:
            oc.append((f"Patta - {t}", v, c))
    other = float(s[(s.ownership == "Patta") & ~lt.isin(["Dry", "Wet"])].site_acre.sum())
    if other > 0:
        oc.append(("Patta - Other", other, "#ff9800"))
    oc.append(("Government", float(gov), "#2ecc71"))
    if unk > 0:
        oc.append(("No A-register record", unk, "#9e9e9e"))
    crops = []
    try:
        cg = pd.read_excel(site.master, sheet_name="Crops_Grown")
        if len(cg):
            eff = cg.groupby("crop_name").agg(acre=("crop_area_ac", "max"), seasons=("season", lambda x: ", ".join(sorted(set(x))))).sort_values("acre", ascending=False)
            crops = [{"crop": k, "acre": float(r.acre), "seasons": r.seasons} for k, r in eff.iterrows()]
    except ValueError:
        pass
    return {"total": A, "dc_glv": _dc_glv(s, pa, summ), "own": {"Government Land": row(gov), "Patta Land": row(pat), **({"No A-register record": row(unk)} if unk > 0 else {})},
            "unknown_total": unk,
            "lclass": {names.get(k, k): row(v) for k, v in s.groupby("land_type_english").site_acre.sum().sort_values(ascending=False).items()},
            "rem": {k: {"total": round(float(v), 2), "pct": round(float(v) / gov * 100, 2) if gov else 0} for k, v in rem.items()},
            "gov_total": float(gov), "patta_total": float(pat),
            "glv_unit": (glv_final(site) or {}).get("units") or {"acres": {"range": ac_r, "area": ac_a}, "sqft": {"range": sq_r, "area": sq_a}, "nodata": {"range": "–", "area": max(nod, 0.0)}, "total": float(pat)},
            "glv_area": (glv_final(site) or {}).get("C1", {}).get("total", {}).get("acre") or A, "glv_source": lc.get("final_source"),
            "c1": cond("C1"), "c2": cond("C2"), "villages": {k: {"acre": float(v), "pct": float(v / A * 100)} for k, v in vil.items()},
            "crop": {"crop": crop["crop_acre"], "nosurvey": crop["no_survey_acre"], "noncrop": crop["noncrop_acre"]},
            "village_colors": summ.get("village_colors", {}), "split_surveys": int(s[s.split.astype(str).str.lower() == "true"].survey_uid.nunique()),   # cut surveys still in the site after the edge review
            "abstract": abstract, "own_classes": oc, "crops": crops if summ.get("crop") else []}


def compose_streetview(raw, out, AC, pins, info):
    """street-view map (pins on the road render) + one full-slide picture per pin that has a photo. Returns pin hotspots (px)."""
    e = info["road"]
    o = Image.new("RGB", (W, H + B), (16, 58, 78)); o.paste(Image.open(raw / "road_connectivity_raw.png").convert("RGB"), (0, B)); d = ImageDraw.Draw(o)
    t = "STREET VIEW"; f = F(66); d.text(((W - d.textlength(t, font=f)) / 2, (B - 76) / 2), t, fill=(255, 235, 0), font=f)
    to3857 = lambda lon, lat: (lon * 20037508.34 / 180, math.log(math.tan((90 + lat) * math.pi / 360)) * 6378137)
    pxe = lambda x, y: ((x - e[0]) / (e[2] - e[0]) * W, B + (e[3] - y) / (e[3] - e[1]) * H)
    hot = []
    for p in pins:
        x, y = pxe(*to3857(p["lon"], p["lat"]))
        pin, (tx, ty) = teardrop(r=34, h=118, col=(16, 90, 200)); o.paste(pin, (int(round(x - tx)), int(round(y - ty))), pin)
        cy = y - 118 + 34 + 3; fn = F(34); n = str(p["n"])
        d.ellipse([x - 24, cy - 24, x + 24, cy + 24], fill=(255, 255, 255)); d.text((x - d.textlength(n, font=fn) / 2, cy - 20), n, fill=(16, 58, 78), font=fn)
        hot.append({"n": p["n"], "x": x - 40, "y": y - 124, "w": 80, "h": 130})
    qx, qy = pxe(*info["site_pole"]); lt = f"Site Boundary ({AC} Acres)"
    if any(abs(h["x"] + 40 - qx) < 260 and abs(h["y"] + 124 - qy) < 90 for h in hot):
        qy += 80   # keep the label clear of a pin
    d.text((qx - d.textlength(lt, font=F(36)) / 2, qy - 22), lt, fill=(255, 255, 255), font=F(36), stroke_width=4, stroke_fill=(0, 0, 0))
    lines = [f"{p['n']}. {p['label']}" + (f" ({p['date']})" if p.get("date") else "") for p in pins]
    any_photo = any(p.get("photo") for p in pins)
    lines.append("Click a pin to view " + ("the site photo" if any_photo else "the spot in Google Street View"))
    fp = F(28); pw = max(d.textlength(s_, font=fp) for s_ in lines) + 48; ph = len(lines) * 42 + 70
    corners = [(W - pw - 24, B + 24), (24, B + 24), (24, B + H - ph - 24), (W - pw - 24, B + H - ph - 24)]
    clash = lambda c: sum(1 for h in hot if c[0] - 40 < h["x"] + 40 < c[0] + pw + 40 and c[1] - 20 < h["y"] + 124 < c[1] + ph + 140)
    x0, y0 = min(corners, key=clash)      # the corner with no pin under the legend
    d.rectangle([x0, y0, x0 + pw, y0 + ph], fill=(16, 58, 78), outline=(255, 255, 255), width=2)
    d.text((x0 + 24, y0 + 12), "Street view points", fill=(255, 235, 0), font=F(32))
    for k, s_ in enumerate(lines):
        d.text((x0 + 24, y0 + 58 + k * 42), s_, fill=(255, 255, 255) if k < len(pins) else (255, 235, 0), font=fp)
    o.save(out / "Street_View.png")
    for p in pins:
        if not p.get("photo"):
            continue
        ph_ = Image.open(p["photo"]).convert("RGB"); sl = Image.new("RGB", (W, H + B), (0, 0, 0))
        sc = min(W / ph_.width, H / ph_.height); im = ph_.resize((int(ph_.width * sc), int(ph_.height * sc)), Image.LANCZOS)
        sl.paste(im, ((W - im.width) // 2, B + (H - im.height) // 2)); dd = ImageDraw.Draw(sl)
        dd.rectangle([0, 0, W, B], fill=(16, 58, 78))
        tt = f"STREET VIEW - PIN {p['n']}: {p['label']}"; dd.text(((W - dd.textlength(tt, font=F(52))) / 2, (B - 60) / 2), tt, fill=(255, 235, 0), font=F(52))
        src = ("Imagery: Google Street View" + (f", {p['date']}" if p.get("date") else "")) if p.get("photo_source") == "Google Street View" else ""
        if src:
            dd.text((24, B + H - 52), src, fill=(255, 255, 255), font=F(28), stroke_width=3, stroke_fill=(0, 0, 0))
        dd.text((W - 330, B + H - 52), "Click to go back", fill=(255, 235, 0), font=F(28), stroke_width=3, stroke_fill=(0, 0, 0))
        sl.save(out / f"Street_View_Pin_{p['n']}.png")
    return hot


ICON = {"Seaport": (0, 145, 234), "Airport": (124, 77, 255), "SIPCOT park": (255, 109, 0)}


def compose_regional(raw, out, AC, rg, cn):
    R_ = json.loads((raw / "_region_wide.json").read_text()); e = R_["ext"]
    im = Image.open(raw / "region_wide.png").convert("RGB"); o = Image.new("RGB", (W, H + B), (16, 58, 78)); o.paste(im, (0, B)); d = ImageDraw.Draw(o)
    t = "REGIONAL CONNECTIVITY"; d.text(((W - d.textlength(t, font=F(66))) / 2, (B - 76) / 2), t, fill=(255, 235, 0), font=F(66))
    pxe = lambda x, y: ((x - e[0]) / (e[2] - e[0]) * W, B + (e[3] - y) / (e[3] - e[1]) * H)
    for n, (x, y) in R_["dlab"].items():
        px, py = pxe(x, y); tt = n.upper(); hw = d.textlength(tt, font=F(34)) / 2
        d.text((min(max(px - hw, 10), W - 2 * hw - 10), py), tt, fill=(120, 210, 255), font=F(34), stroke_width=3, stroke_fill=(0, 0, 0))
    bnd = json.loads((raw.parent / "render_params.json").read_text(encoding="utf-8"))
    import geopandas as gpd
    c = gpd.read_file(bnd["boundary"]).geometry.iloc[0].centroid
    sx, sy = pxe(*to3857(c.x, c.y))
    for m in rg["marks"]:
        x, y = pxe(*to3857(m["lon"], m["lat"]))
        inside = 0 <= x <= W and B <= y <= B + H
        if not inside:
            continue
        col = ICON.get(m["kind"], (255, 255, 255))
        d.line([(sx, sy), (x, y)], fill=col, width=5)
        d.ellipse([x - 16, y - 16, x + 16, y + 16], fill=col, outline=(255, 255, 255), width=4)
        lab = m["name"] + f" - {m['road_km'] or m['km']} km"
        tw = d.textlength(lab, font=F(30))
        d.text((x + 22 if x + 22 + tw < W - 10 else x - 22 - tw, y - 18), lab, fill=(255, 255, 255), font=F(30), stroke_width=4, stroke_fill=(0, 0, 0))
    pin, (tx, ty) = teardrop(r=34, h=118); o.paste(pin, (int(round(sx - tx)), int(round(sy - ty))), pin)
    d = ImageDraw.Draw(o); lab = f"Proposed Site ({AC} Acres)"
    d.text((sx - d.textlength(lab, font=F(36)) / 2, sy - 170), lab, fill=(255, 235, 0), font=F(36), stroke_width=4, stroke_fill=(0, 0, 0))
    lines = []
    for m in rg["marks"]:
        rd = f"road {m['road_km']} km, about {m['minutes']} min" if m.get("road_km") else f"{m['km']} km straight line"
        lines.append(f"{m['kind']}: {m['name']} - {rd}")
    if cn.get("nh"):
        lines.append(f"Highway: {cn['nh']}")
    fp = F(27); pw = max(d.textlength(s_, font=fp) for s_ in lines) + 48; ph = len(lines) * 43 + 76
    corners = [(24, B + H - ph - 24), (W - pw - 24, B + H - ph - 24), (24, B + 24), (W - pw - 24, B + 24)]
    x0, y0 = max(corners, key=lambda q: abs(q[0] + pw / 2 - sx) + abs(q[1] + ph / 2 - sy))   # the corner farthest from the site
    d.rectangle([x0, y0, x0 + pw, y0 + ph], fill=(16, 58, 78), outline=(255, 255, 255), width=3); d.text((x0 + 24, y0 + 14), "Key logistics nodes", fill=(255, 235, 0), font=F(34))
    for k, s_ in enumerate(lines):
        d.text((x0 + 24, y0 + 64 + k * 43), s_, fill=(255, 255, 255), font=fp)
    lx, ly = 24, B + 24
    if x0 < 600 and y0 < B + 300:
        lx = W - 430
    items = [((255, 214, 0), "National highway / expressway", "line"), ((63, 182, 255), "District boundary", "line")] + [(v, k, "dot") for k, v in ICON.items()]
    d.rectangle([lx, ly, lx + 450, ly + 40 * len(items) + 24], fill=(16, 58, 78), outline=(255, 255, 255), width=2)
    for k, (col, tx_, kind_) in enumerate(items):
        yy = ly + 22 + k * 40
        if kind_ == "line":
            d.line([(lx + 16, yy + 10), (lx + 64, yy + 10)], fill=col, width=6)
        else:
            d.ellipse([lx + 30, yy, lx + 50, yy + 20], fill=col, outline=(255, 255, 255), width=2)
        d.text((lx + 78, yy - 4), tx_, fill=(255, 255, 255), font=F(25))
    o.save(out / "Regional_Context.png")


def compose_photos(raw, out, AC, photos):
    I = json.loads((raw / "_render_info.json").read_text()); e = I["boundary"]
    o = Image.new("RGB", (W, H + B), (16, 58, 78)); o.paste(Image.open(raw / "boundary_raw.png").convert("RGB"), (0, B)); d = ImageDraw.Draw(o)
    t = "SITE PHOTOS"; d.text(((W - d.textlength(t, font=F(66))) / 2, (B - 76) / 2), t, fill=(255, 235, 0), font=F(66))
    pxe = lambda x, y: ((x - e[0]) / (e[2] - e[0]) * W, B + (e[3] - y) / (e[3] - e[1]) * H)
    n = len(photos); TW, TH = 420, 300
    slots = [(30, B + 30), (W - TW - 30, B + 30), (30, B + H - TH - 30), (W - TW - 30, B + H - TH - 30),
             (30, B + (H - TH) // 2), (W - TW - 30, B + (H - TH) // 2), ((W - TW) // 2, B + 30), ((W - TW) // 2, B + H - TH - 30)]
    pts = [pxe(*to3857(p["pos"][1], p["pos"][0])) if p.get("pos") else None for p in photos]
    free = list(range(min(n, len(slots))))
    order = []
    for i, p in enumerate(pts):       # nearest free slot to each photo's spot
        if p is None:
            continue
        k = min(free, key=lambda j: math.dist((slots[j][0] + TW / 2, slots[j][1] + TH / 2), p)) if free else None
        if k is not None:
            free.remove(k); order.append((i, k))
    for i, p in enumerate(pts):
        if p is None and free:
            order.append((i, free.pop(0)))
    for i, k in order:
        x0, y0 = slots[k]
        if pts[i]:
            d.line([(x0 + TW / 2, y0 + TH / 2), pts[i]], fill=(255, 255, 255), width=5)
            d.ellipse([pts[i][0] - 14, pts[i][1] - 14, pts[i][0] + 14, pts[i][1] + 14], fill=(255, 23, 68), outline=(255, 255, 255), width=4)
    d.text((W / 2 - 200, B + H / 2), f"{AC} Acres", fill=(255, 235, 0), font=F(54), stroke_width=4, stroke_fill=(0, 0, 0))
    for i, k in order:
        x0, y0 = slots[k]
        th = Image.open(photos[i]["file"]).convert("RGB"); th.thumbnail((TW, TH))
        box = Image.new("RGB", (TW + 12, TH + 12), (255, 255, 255)); box.paste(th, (6 + (TW - th.width) // 2, 6 + (TH - th.height) // 2))
        o.paste(box, (x0 - 6, y0 - 6)); d = ImageDraw.Draw(o)
        d.rectangle([x0, y0, x0 + 58, y0 + 52], fill=(16, 58, 78)); d.text((x0 + 16, y0 + 6), str(photos[i]["n"]), fill=(255, 235, 0), font=F(38))
    o.save(out / "Site_Photos.png")


def build_ppt(site, img, AC, summ, keep, pins=None, hot=None):
    """keep: slide keys to include - only slides whose stage was ticked and finished (title and land summary always)."""
    from pptx import Presentation
    from pptx.util import Emu
    tpl = SETTINGS.ppt_template
    if not tpl.exists():
        raise RuntimeError(f"PPT template not found: {tpl} (set PPT_TEMPLATE in .env)")
    shutil.copy(tpl, site.ppt)
    prs = Presentation(site.ppt)
    # template = Perungalathur v2 layout: 1 title, 2 location, 3 boundary, 4 road, 5 street view, 6-10 summaries, 11 DC land, 12-15 hidden photos
    key_of = {2: "location", 3: "boundary", 4: "road", 5: "streetview", 6: "land", 7: "glv", 8: "village", 9: "ownership", 10: "crop", 11: "dc"}
    imgs = {2: "Location_Map", 3: "Boundary", 4: "Road_Connectivity", 5: "Street_View", 6: "s06_land_summary", 7: "s07_glv_summary",
            8: "s08_villagewise", 9: "s09_ownership", 10: "s10_crop"}
    slides = list(prs.slides)
    for n, name in imgs.items():
        if key_of[n] not in keep:
            continue
        sl = slides[n - 1]
        pics = [sh for sh in sl.shapes if sh.shape_type == 13 and sh.width > 0.9 * prs.slide_width]
        sl.part.related_part(pics[0]._element.blipFill.blip.rEmbed)._blob = (img / f"{name}.png").read_bytes()
    drop = [n for n, k in key_of.items() if k not in keep]      # template slide numbers (1-based) to remove at the end
    photo_slides = slides[11:15]
    ids_ = prs.slides._sldIdLst
    tmpl = list(ids_)                                              # template slide id elements, positions fixed before inserting
    extra = 0
    last_after = {}
    # new slides: government land after ownership (template 9); soil and sensitive land after crop (template 10)
    for k_, name, anchor in (("summary", "s01b_summary", 1), ("regional", "Regional_Context", 2), ("sitephotos", "Site_Photos", 5), ("glvmap", "s07b_glvmap", 7), ("govtland", "s09b_govt", 9), ("owners", "s09c_owners", 9), ("companies", "s09d_company", 9),
                             ("soil", "s11_soil", 10), ("sensitive", "s12_sensitive", 10)):
        if k_ not in keep:
            continue
        ns = prs.slides.add_slide(slides[9].slide_layout)
        for ph in list(ns.placeholders):
            ph._element.getparent().remove(ph._element)
        ns.shapes.add_picture(str(img / f"{name}.png"), 0, 0, prs.slide_width, prs.slide_height)
        el = ids_[-1]; ids_.remove(el)
        after = last_after.get(anchor, tmpl[anchor - 1])
        ids_.insert(list(ids_).index(after) + 1, el); last_after[anchor] = el
    if "streetview" in keep and pins:
        sv = slides[4]; k = prs.slide_width / W
        by_n = {h["n"]: h for h in hot}
        for i in range(1, 5):
            shp = [x for x in sv.shapes if x.name == f"Pin_{i}"]
            p = next((q for q in pins if q["n"] == i), None)
            if not shp:
                continue
            shp = shp[0]
            if p is None:
                shp._element.getparent().remove(shp._element); continue
            h = by_n[i]
            shp.left, shp.top, shp.width, shp.height = Emu(int(h["x"] * k)), Emu(int(h["y"] * k)), Emu(int(h["w"] * k)), Emu(int(h["h"] * k))
            if p.get("photo"):
                ps = photo_slides[i - 1]
                pic = [x for x in ps.shapes if x.shape_type == 13][0]
                ps.part.related_part(pic._element.blipFill.blip.rEmbed)._blob = (img / f"Street_View_Pin_{i}.png").read_bytes()
                shp.click_action.target_slide = ps
                for x in ps.shapes:
                    if x.name in ("Back_to_Street_View", "Close"):
                        x.click_action.target_slide = sv
            else:
                shp.click_action.hyperlink.address = p["url"]
        drop += [12 + i - 1 for i in range(1, 5) if not next((q for q in pins if q["n"] == i and q.get("photo")), None)]
    else:
        drop += [12, 13, 14, 15]
    vs = [v["village"] for v in summ["villages"]]; tl = list(dict.fromkeys(v["taluk"] for v in summ["villages"])); ds = list(dict.fromkeys(v["district"] for v in summ["villages"]))
    tb = [sh for sh in slides[0].shapes if sh.has_text_frame and "Preliminary" in sh.text_frame.text][0]
    p0, p2 = tb.text_frame.paragraphs[0], tb.text_frame.paragraphs[2]
    r = p0.runs
    r[1].text = f"Preliminary Study of the Proposed Site Covering {AC} Acres in "; r[2].text = _and(tl); r[3].text = f" Taluk{'s' if len(tl) > 1 else ''}, "
    r[4].text = _and(ds); r[5].text = f" District{'s' if len(ds) > 1 else ''}"
    r = p2.runs
    for x in r[2:]:
        x.text = ""
    r[2].text = _and(vs)
    geom_c = summ.get("centroid")
    tb3 = [sh for sh in slides[2].shapes if sh.has_text_frame and "Google Maps" in sh.text_frame.text]
    if tb3 and geom_c:
        url = f"https://www.google.com/maps?q={geom_c[1]:.6f},{geom_c[0]:.6f}&t=k"
        run_ = [x for x in tb3[0].text_frame.paragraphs[0].runs if "google.com" in x.text][0]; run_.text = url; run_.hyperlink.address = url
    ids = prs.slides._sldIdLst
    for n in sorted(set(drop)):
        if n - 1 < len(tmpl):
            sid = tmpl[n - 1]; prs.part.drop_rel(sid.rId); ids.remove(sid)
    prs.core_properties.title = f"Preliminary Site Analysis - {site.name} ({AC} Acres)"
    prs.save(site.ppt)
    return len(prs.slides._sldIdLst)


OPTIONAL = ("maps", "landcost", "crop", "streetview", "soil", "sensitive", "govtland", "glvmap", "owners", "companies", "dimensions", "regional", "sitephotos", "summary")


def _included(site, ctx):
    st = db.stage_status(site.site_id)
    return sorted(x for x in OPTIONAL if ctx.is_on(x) and st.get(x, ("",))[0] == "done")


def stale(site, ctx):
    """re-build the PPT when the ticked optional stages differ from the last build (e.g. crop unticked -> no crop slide)"""
    return db.get_summary(site.site_id).get("ppt_included") != _included(site, ctx) or not site.ppt.exists()


def exec_data(summ, data, on):
    A = data["total"]; m = summ.get("master", {}); vs = summ.get("villages", [])
    cards = [{"k": "Site area (A-register)", "v": f"{A:,.2f} <small>acres</small>", "s": f"{summ.get('surveys', 0)} surveys · {m.get('subdivisions', 0)} subdivisions · {len(vs)} villages"},
             {"k": "Patta / Government", "v": f"{data['patta_total'] / A * 100:.1f}% <small>patta</small>", "s": f"Patta {data['patta_total']:,.2f} ac · Govt {data['gov_total']:,.2f} ac"},
             {"k": "Location", "v": ", ".join(dict.fromkeys(v["district"] for v in vs)), "s": "Taluk: " + ", ".join(dict.fromkeys(v["taluk"] for v in vs))}]
    lc = summ.get("landcost", {}).get("C1") if on("landcost") else None
    if lc:
        c2 = summ["landcost"]["C2"]
        cards.append({"k": "Estimated cost per acre", "v": f"{lc['act_per_acre'] / 1e5:.2f} <small>lakh (Act, C1)</small>", "cls": "o",
                      "s": f"Negotiation {lc['neg_per_acre'] / 1e5:.2f} L · C2: {c2['act_per_acre'] / 1e5:.2f} / {c2['neg_per_acre'] / 1e5:.2f} L"})
    cr = summ.get("crop") if on("crop") else None
    if cr:
        cards.append({"k": "Crop land", "v": f"{cr['crop_acre'] / A * 100:.1f}% <small>of site</small>", "cls": "g",
                      "s": f"Crop {cr['crop_acre']:,.2f} ac · non-crop {cr['noncrop_acre']:,.2f} · no survey {cr['no_survey_acre']:,.2f}"})
    if lc and summ["landcost"].get("final_source"):   # final GLV calculation supplied: show its land cost, not this app's GLV total
        cards.append({"k": "Land cost – final GLV (C1)", "v": f"Rs {lc['act_total'] / 1e7:,.2f} <small>crore (Act)</small>",
                      "s": f"Negotiation Rs {lc['neg_total'] / 1e7:,.2f} cr · {lc['total_acre']:,.2f} ac"})
    elif m.get("glv_crore") and on("glv"):
        cards.append({"k": "Guideline value (indicative)", "v": f"Rs {m['glv_crore']:,.1f} <small>crore</small>", "s": "TNReginet rates as quoted (acre / sq ft)"})
    se = summ.get("sensitive") if on("sensitive") else None
    if se:
        cards.append({"k": "Sensitive land inside", "v": f"{se['sensitive_acre']:,.2f} <small>acres</small>", "cls": "o",
                      "s": ", ".join(f"{r['category']} {r['acre']:.1f}" for r in se["rows"][:3])})
    ow = summ.get("owners") if on("owners") else None
    if ow:
        cards.append({"k": "Patta holdings", "v": f"{ow['pattas']} <small>pattas</small>", "s": f"{ow['companies']} held by companies · {ow['owner_entries']} owner entries"})
    cn = summ.get("connectivity", {})
    notes = (cn.get("lines") or [])[:4] + ([cn["sipcot_park"]] if cn.get("sipcot_park") else [])
    return {"cards": cards[:8], "notes": notes if on("maps") else []}


ZERO_COST = {x: 0.0 for x in ("gov_acre", "pat_acre", "act_gov", "neg_gov", "act_pat", "neg_pat", "act_total", "neg_total", "act_per_acre", "neg_per_acre")}


def run(site, ctx):
    summ = db.get_summary(site.site_id)
    import geopandas as gpd
    c = gpd.read_file(site.boundary).geometry.iloc[0].centroid
    db.save_summary(site.site_id, {"centroid": [c.x, c.y]}); summ["centroid"] = [c.x, c.y]
    st = db.stage_status(site.site_id)
    on = lambda s_: ctx.is_on(s_) and st.get(s_, ("",))[0] == "done"     # ticked for this run AND finished
    out = site.maps_dir; raw = site.work / "render"
    has_maps = on("maps") and (raw / "_render_info.json").exists()
    has_cost = on("landcost") and "landcost" in summ
    has_crop = on("crop") and "crop" in summ
    has_sv = has_maps and on("streetview") and (site.work / "streetview.json").exists()
    has_soil = has_maps and on("soil") and "soil" in summ and (raw / "map_soil.png").exists()
    has_sens = has_maps and on("sensitive") and "sensitive" in summ and (raw / "map_sensitive.png").exists()
    has_govt = has_maps and on("govtland") and "govtland" in summ and (raw / "map_govt.png").exists()
    has_glvmap = has_maps and on("glvmap") and "glvmap" in summ and (raw / "map_glv.png").exists()
    has_owners = has_maps and on("owners") and "owners" in summ and (raw / "map_holding.png").exists()
    has_summary = on("summary")
    has_regional = has_maps and on("regional") and (raw / "region_wide.png").exists() and "regional" in summ
    has_photos = has_maps and on("sitephotos") and bool((summ.get("sitephotos") or {}).get("photos"))
    has_comp = has_maps and on("companies") and (summ.get("companies") or {}).get("n", 0) > 0 and (raw / "map_company.png").exists()
    if not (on("dimensions") and "dimensions" in summ):
        summ = {**summ, "dimensions": None}
    keep = {"land"} | ({"location", "boundary", "road", "village", "ownership"} if has_maps else set()) | ({"glv"} if has_cost else set()) \
        | ({"crop"} if has_crop and has_maps else set()) | ({"streetview"} if has_sv else set())         | ({"soil"} if has_soil else set()) | ({"sensitive"} if has_sens else set()) | ({"govtland"} if has_govt else set()) \
        | ({"glvmap"} if has_glvmap else set()) | ({"owners"} if has_owners else set()) | ({"summary"} if has_summary else set()) | ({"companies"} if has_comp else set()) \
        | ({"regional"} if has_regional else set()) | ({"sitephotos"} if has_photos else set())
    if not has_cost:
        summ = {**summ, "landcost": {k: {**ZERO_COST, "total_acre": 1.0} for k in ("C1", "C2")}}
    if not has_crop:
        summ = {**summ, "crop": None}
    data = slide_data(site, summ)
    AC = f"{data['total']:.2f}"
    if has_soil:
        data["soil"] = summ["soil"]
        sc_ = raw / "_soil_colors.json"
        data["tharam_colors"] = json.loads(sc_.read_text())["tharam_colors"] if sc_.exists() else {}
    if has_sens:
        data["sensitive"] = summ["sensitive"]
    if has_govt:
        data["govtland"] = summ["govtland"]
    if has_glvmap:
        data["glvmap"] = summ["glvmap"]
    if has_owners:
        data["owners"] = summ["owners"]
    if has_comp:
        data["companies"] = summ["companies"]
    if has_summary:
        data["exec"] = exec_data(summ, data, on)
    pins = hot = None
    if has_maps:
        for f in ("map_village.png", "map_owner.png", "map_crop.png", "map_soil.png", "map_sensitive.png", "map_govt.png", "map_glv.png", "map_holding.png", "map_company.png"):
            if (raw / f).exists():
                shutil.copy(raw / f, out / f)
        cn = json.loads((site.work / "connectivity.json").read_text(encoding="utf-8"))
        compose(site, raw, out, AC, summ, cn)
        if has_regional:
            compose_regional(raw, out, AC, summ["regional"], cn)
        if has_photos:
            compose_photos(raw, out, AC, summ["sitephotos"]["photos"])
        if has_sv:
            pins = json.loads((site.work / "streetview.json").read_text(encoding="utf-8"))
            hot = compose_streetview(raw, out, AC, pins, json.loads((raw / "_render_info.json").read_text()))
    (out / "data.json").write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
    ctx.log("building summary slides (headless Chrome) ...")
    p = subprocess.run([sys.executable, "-m", "sitesuit.slides_html", str(out)], cwd=str(Path(__file__).resolve().parents[2]), capture_output=True,
                       text=True, encoding="utf-8", errors="replace", timeout=600)
    if p.returncode != 0:
        raise RuntimeError("slide build failed: " + (p.stderr or p.stdout)[-800:])
    for f in list(out.glob("s*_*.html")) + [out / "data.json"]:
        f.unlink(missing_ok=True)
    n = build_ppt(site, out, AC, summ, keep, pins, hot)
    db.register_document(site.site_id, "ppt", site.ppt)
    db.save_summary(site.site_id, {"ppt_included": _included(site, ctx)})
    left = [k for k, s_ in (("location / boundary / road / village / ownership slides", "maps"), ("GLV + land cost slide", "landcost"),
                            ("crop slide", "crop"), ("street view", "streetview"), ("soil slide", "soil"), ("sensitive land slide", "sensitive"), ("government land slide", "govtland"), ("GLV map", "glvmap"), ("patta holdings", "owners"), ("company ownership", "companies"),
                            ("boundary dimensions", "dimensions"), ("regional context", "regional"), ("site photos", "sitephotos"),
                            ("executive summary", "summary")) if not on(s_)]
    return f"{site.ppt.name}: {n} slides, {AC} acres" + (f"; left out: {', '.join(left)}" if left else "") + " (PDF only after you confirm the slides)"


def export_pdf(site):
    """PowerPoint COM export - only when the user confirms the PPT."""
    pdf = site.ppt.with_suffix(".pdf")
    ps = (f"$pp = New-Object -ComObject PowerPoint.Application; try {{ $p = $pp.Presentations.Open('{site.ppt}', $true, $false, $false); "
          f"$p.SaveAs('{pdf}', 32); $p.Close() }} finally {{ $pp.Quit() }}")
    subprocess.run(["powershell", "-NoProfile", "-Command", ps], check=True, timeout=300)
    db.register_document(site.site_id, "pdf", pdf)
    return pdf
